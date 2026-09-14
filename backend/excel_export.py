"""Xuất bộ test case ra Excel:
- Sheet "Tổng hợp" (overview): kết quả test tổng + thống kê theo màn hình + biểu đồ.
- Mỗi MÀN HÌNH một sheet riêng, case trong sheet gom theo Chức năng, đủ 14 cột.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# (khóa dữ liệu, tiêu đề cột, độ rộng). "stt" là số thứ tự tự đánh.
COLUMNS = [
    ("stt",          "STT",                     6),
    ("id",           "Mã TC",                   12),
    ("account_type", "Loại tài khoản",          16),
    ("screen",       "Màn hình",                18),
    ("function",     "Chức năng",               22),
    ("sub_function", "Chức năng con",           20),
    ("ui_function",  "UI/Function",             12),
    ("precondition", "Tiền điều kiện",          24),
    ("steps",        "Bước tái hiện",           46),
    ("expected",     "Kết quả mong muốn",       34),
    ("actual",       "Kết quả thực tế",         26),
    ("status",       "Tình trạng",              13),
    ("bug_link",     "Dẫn link bug (Clickup)",  22),
    ("test_date",    "Ngày test",               14),
]

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
TITLE_FILL = PatternFill("solid", fgColor="1F3864")
GROUP_FILL = PatternFill("solid", fgColor="DDEBF7")     # dòng tiêu đề nhóm chức năng
ROW_ALT_FILL = PatternFill("solid", fgColor="EAF1FB")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
BLOCK_FILL = PatternFill("solid", fgColor="FFEB9C")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")

THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
GROUP_FONT = Font(name="Calibri", size=10, bold=True, color="1F3864")
CELL_FONT = Font(name="Calibri", size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _classify_status(value: str) -> str:
    v = str(value).strip().lower()
    if v in ("pass", "passed", "đạt", "ok"):
        return "pass"
    if v in ("fail", "failed", "lỗi", "không đạt", "ng"):
        return "fail"
    if v in ("blocked", "block", "pending"):
        return "blocked"
    return "untested"


def _safe_sheet_name(name: str, used: set) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", str(name)).strip() or "Màn hình"
    name = name[:28]
    base, i = name, 2
    while name.lower() in used:
        name = f"{base[:25]} {i}"
        i += 1
    used.add(name.lower())
    return name


def build_workbook(test_cases: list[dict], url: str = "") -> Workbook:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Tổng hợp"

    # Gom test case theo Màn hình (giữ thứ tự xuất hiện)
    by_screen, order = {}, []
    for c in test_cases:
        scr = c.get("screen") or "Chung"
        if scr not in by_screen:
            by_screen[scr] = []
            order.append(scr)
        by_screen[scr].append(c)

    _build_summary_sheet(summary_ws, test_cases, by_screen, order, url)

    # Mỗi màn hình một sheet, case gom theo Chức năng nhưng GIỮ THỨ TỰ trên→dưới
    used = set()
    for scr in order:
        ws = wb.create_sheet(_safe_sheet_name(scr, used))
        _fill_testcases_sheet(ws, _group_by_function_ordered(by_screen[scr]), scr)
    return wb


def _group_by_function_ordered(cases: list) -> list:
    """Gom case theo Chức năng nhưng giữ thứ tự xuất hiện (đọc từ trên xuống),
    KHÔNG sắp xếp theo bảng chữ cái — để đúng thứ tự test trên màn hình."""
    groups, order = {}, []
    for c in cases:
        fn = c.get("function") or "Chung"
        if fn not in groups:
            groups[fn] = []
            order.append(fn)
        groups[fn].append(c)
    out = []
    for fn in order:
        out.extend(groups[fn])
    return out


def _kpi_card(ws, anchor_row, col, label, value, fill_hex):
    """Vẽ 1 thẻ KPI (gộp 2x2 ô): nhãn nhỏ trên, số to dưới."""
    c1, c2 = col, col + 1
    fill = PatternFill("solid", fgColor=fill_hex)
    ws.merge_cells(start_row=anchor_row, start_column=c1, end_row=anchor_row, end_column=c2)
    lab = ws.cell(row=anchor_row, column=c1, value=label)
    lab.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    lab.alignment = CENTER; lab.fill = fill
    ws.merge_cells(start_row=anchor_row + 1, start_column=c1, end_row=anchor_row + 2, end_column=c2)
    val = ws.cell(row=anchor_row + 1, column=c1, value=value)
    val.font = Font(name="Calibri", size=22, bold=True, color="FFFFFF")
    val.alignment = CENTER; val.fill = fill
    for rr in (anchor_row, anchor_row + 1, anchor_row + 2):
        for cc in (c1, c2):
            ws.cell(row=rr, column=cc).fill = fill


def _build_summary_sheet(ws, test_cases, by_screen, order, url):
    last_col = "G"
    total = len(test_cases)
    cnt = {"pass": 0, "fail": 0, "blocked": 0, "untested": 0}
    for c in test_cases:
        cnt[_classify_status(c.get("status", ""))] += 1
    tpct = round(cnt["pass"] / total * 100) if total else 0
    done = cnt["pass"] + cnt["fail"] + cnt["blocked"]
    dpct = round(done / total * 100) if total else 0

    ws.merge_cells(f"A1:{last_col}1")
    b = ws.cell(row=1, column=1, value="BÁO CÁO TỔNG HỢP KẾT QUẢ TEST")
    b.font = TITLE_FONT; b.fill = TITLE_FILL; b.alignment = CENTER
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{last_col}2")
    info = ws.cell(row=2, column=1, value=f"URL: {url}")
    info.font = Font(name="Calibri", size=10, italic=True, color="44546A")
    info.alignment = Alignment(horizontal="left", vertical="center")

    # --- Dải KPI ---
    _kpi_card(ws, 4, 1, "TỔNG TEST CASE", total, "2F5597")
    _kpi_card(ws, 4, 3, "SỐ MÀN HÌNH", len(order), "548235")
    _kpi_card(ws, 4, 5, "TỶ LỆ PASS", f"{tpct}%", "C55A11")
    ws.merge_cells("G4:G4")
    prog = ws.cell(row=4, column=7, value="Đã test")
    prog.font = Font(name="Calibri", size=10, bold=True, color="1F3864"); prog.alignment = CENTER
    ws.merge_cells("G5:G6")
    pv = ws.cell(row=5, column=7, value=f"{dpct}%")
    pv.font = Font(name="Calibri", size=18, bold=True, color="1F3864"); pv.alignment = CENTER
    for rr in (4, 5, 6):
        ws.row_dimensions[rr].height = 22

    # --- Bảng kết quả theo trạng thái ---
    hr = 8
    ws.cell(row=hr - 1, column=1, value="KẾT QUẢ THEO TÌNH TRẠNG").font = Font(
        name="Calibri", size=11, bold=True, color="1F3864")
    for col, label in enumerate(["Tình trạng", "Số lượng", "Tỷ lệ"], start=1):
        cell = ws.cell(row=hr, column=col, value=label)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = CENTER; cell.border = BORDER
    rows = [("Pass (Đạt)", cnt["pass"], PASS_FILL), ("Fail (Không đạt)", cnt["fail"], FAIL_FILL),
            ("Blocked", cnt["blocked"], BLOCK_FILL), ("Chưa test", cnt["untested"], None)]
    r = hr + 1
    chart_first = r
    for label, num, fill in rows:
        pct = round(num / total * 100) if total else 0
        c1 = ws.cell(row=r, column=1, value=label)
        c2 = ws.cell(row=r, column=2, value=num)
        c3 = ws.cell(row=r, column=3, value=f"{pct}%")
        for c in (c1, c2, c3):
            c.font = CELL_FONT; c.border = BORDER
        c1.alignment = WRAP_TOP; c2.alignment = CENTER; c3.alignment = CENTER
        if fill:
            c1.fill = fill
        r += 1
    chart_last = r - 1
    ct = ws.cell(row=r, column=1, value="TỔNG CỘNG")
    cn = ws.cell(row=r, column=2, value=total)
    cp = ws.cell(row=r, column=3, value="100%" if total else "0%")
    for c in (ct, cn, cp):
        c.font = Font(name="Calibri", size=10, bold=True); c.fill = TOTAL_FILL; c.border = BORDER
    ct.alignment = WRAP_TOP; cn.alignment = CENTER; cp.alignment = CENTER

    if total:
        chart = PieChart(); chart.title = "Phân bố kết quả test"
        chart.height = 7.5; chart.width = 11
        data = Reference(ws, min_col=2, min_row=hr, max_row=chart_last)
        cats = Reference(ws, min_col=1, min_row=chart_first, max_row=chart_last)
        chart.add_data(data, titles_from_data=True); chart.set_categories(cats)
        ws.add_chart(chart, "E8")

    # --- Bảng thống kê theo Màn hình (có % Pass tô màu) ---
    sr = r + 3
    ws.cell(row=sr - 1, column=1, value="THỐNG KÊ THEO MÀN HÌNH").font = Font(
        name="Calibri", size=11, bold=True, color="1F3864")
    headers2 = ["Màn hình", "Tổng", "Pass", "Fail", "Blocked", "Chưa test", "% Pass"]
    for col, label in enumerate(headers2, start=1):
        cell = ws.cell(row=sr, column=col, value=label)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = CENTER; cell.border = BORDER
    rr = sr + 1
    for scr in order:
        s = {"total": 0, "pass": 0, "fail": 0, "blocked": 0, "untested": 0}
        for c in by_screen[scr]:
            s["total"] += 1
            s[_classify_status(c.get("status", ""))] += 1
        pct = round(s["pass"] / s["total"] * 100) if s["total"] else 0
        vals = [scr, s["total"], s["pass"], s["fail"], s["blocked"], s["untested"], f"{pct}%"]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=rr, column=col, value=val)
            cell.font = CELL_FONT; cell.border = BORDER
            cell.alignment = WRAP_TOP if col == 1 else CENTER
        pcell = ws.cell(row=rr, column=7)
        if pct >= 80:
            pcell.fill = PASS_FILL
        elif pct >= 50:
            pcell.fill = BLOCK_FILL
        elif s["pass"] or s["fail"]:
            pcell.fill = FAIL_FILL
        rr += 1

    for col, w in enumerate([24, 10, 9, 9, 11, 12, 10], start=1):
        ws.column_dimensions[get_column_letter(col)].width = w


def _fill_testcases_sheet(ws, test_cases: list[dict], screen: str = ""):
    ncols = len(COLUMNS)
    last_col = get_column_letter(ncols)

    ws.merge_cells(f"A1:{last_col}1")
    banner = ws.cell(row=1, column=1, value=f"TEST CASE — {screen}" if screen else "BẢNG TEST CASE")
    banner.font = TITLE_FONT; banner.fill = TITLE_FILL; banner.alignment = CENTER
    ws.row_dimensions[1].height = 26

    header_row = 2
    for col, (_k, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = CENTER; cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[header_row].height = 28

    r = header_row + 1
    stt = 0
    cur_fn = None
    for case in test_cases:
        fn = case.get("function") or "Chung"
        # Dòng tiêu đề nhóm khi đổi chức năng
        if fn != cur_fn:
            cur_fn = fn
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            gcell = ws.cell(row=r, column=1, value=f"▸ Chức năng: {fn}")
            gcell.font = GROUP_FONT; gcell.fill = GROUP_FILL
            gcell.alignment = Alignment(horizontal="left", vertical="center")
            gcell.border = BORDER
            r += 1
        stt += 1
        for col, (key, _label, _w) in enumerate(COLUMNS, start=1):
            value = stt if key == "stt" else case.get(key, "")
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = CELL_FONT; cell.border = BORDER
            cell.alignment = CENTER if key in ("stt", "id", "ui_function", "status", "test_date") else WRAP_TOP
            # (Không tô màu xen kẽ các dòng test case — để dòng trắng sạch)
            if key == "status":
                st = _classify_status(value)
                if st == "pass":
                    cell.fill = PASS_FILL
                elif st == "fail":
                    cell.fill = FAIL_FILL
                elif st == "blocked":
                    cell.fill = BLOCK_FILL
        r += 1

    last = r - 1
    if test_cases:
        status_col = get_column_letter([k for k, *_ in COLUMNS].index("status") + 1)
        dv = DataValidation(type="list", formula1='"Chưa test,Pass,Fail,Blocked"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{status_col}{header_row + 1}:{status_col}{last}")

    ws.freeze_panes = f"C{header_row + 1}"


def to_bytes(test_cases: list[dict], url: str = "") -> bytes:
    wb = build_workbook(test_cases, url)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_workbook(data: bytes) -> list[dict]:
    """Đọc lại file Excel test case (định dạng do tool xuất) -> danh sách case dict.

    Dùng cho chức năng 'gộp vào file có sẵn'. Bỏ qua sheet Tổng hợp, dòng banner,
    dòng tiêu đề cột và dòng tiêu đề nhóm chức năng.
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    label_to_key = {label: key for key, label, _ in COLUMNS}
    cases = []
    for ws in wb.worksheets:
        if ws.title.strip().lower() in ("tổng hợp", "tong hop", "overview"):
            continue
        # Tìm dòng tiêu đề cột (chứa "Mã TC") và lập bản đồ cột -> khóa
        header_row, colmap = None, {}
        for r in range(1, min(ws.max_row, 6) + 1):
            labels = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value else ""
                      for c in range(1, ws.max_column + 1)]
            if "Mã TC" in labels:
                header_row = r
                for c, lbl in enumerate(labels, start=1):
                    if lbl in label_to_key and label_to_key[lbl] != "stt":
                        colmap[c] = label_to_key[lbl]
                break
        if not header_row:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            row = {key: ("" if ws.cell(r, c).value is None else str(ws.cell(r, c).value).strip())
                   for c, key in colmap.items()}
            # Bỏ dòng nhóm/dòng rỗng (không có Mã TC và không có bước)
            if not row.get("id", "") and not row.get("steps", ""):
                continue
            if not row.get("screen", ""):
                row["screen"] = ws.title
            cases.append(row)
    return cases
