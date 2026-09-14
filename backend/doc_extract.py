"""Trích xuất văn bản từ file tài liệu đặc tả: .txt/.md/.csv/.docx/.pdf."""
import io


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith((".txt", ".md", ".csv", ".log")):
        return _decode(data)
    if name.endswith(".docx"):
        return _from_docx(data)
    if name.endswith(".pdf"):
        return _from_pdf(data)
    if name.endswith((".xlsx", ".xlsm")):
        return _from_xlsx(data)
    # Mặc định thử đọc như văn bản thường
    return _decode(data)


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "cp1258", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _from_docx(data: bytes) -> str:
    from docx import Document  # python-docx
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Lấy cả nội dung trong bảng (đặc tả hay để dạng bảng)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n".join(pages)


def _from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
